import pandas as pd
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
import os
from datetime import datetime
from tkinter import Tk, filedialog, messagebox

# --- DICCIONARIO DE CLASIFICACIÓN INTEGRADO ---
# Este diccionario sustituye al archivo CSV de clasificación
CLASIFICACION_INTERNA = {
    'Ayuntamiento': 'Destacadas', 'Congreso': 'Destacadas', 'Delegación': 'Destacadas',
    'IEEH': 'Destacadas', 'Independiente': 'Destacadas', 'MC': 'Destacadas',
    'Morena': 'Destacadas', 'Otros': 'Destacadas', 'PAN': 'Destacadas',
    'Panalh': 'Destacadas', 'PRD': 'Destacadas', 'Presidencia': 'Destacadas',
    'Pri': 'Destacadas', 'PT': 'Destacadas', 'PVEM': 'Destacadas',
    'Bienestar': 'Gobierno del estado', 'CGCG': 'Gobierno del estado',
    'Contraloría': 'Gobierno del estado', 'Cultura': 'Gobierno del estado',
    'Despacho': 'Gobierno del estado', 'DIFH': 'Gobierno del estado',
    'Finanzas': 'Gobierno del estado', 'Gabinete de Seguridad': 'Gobierno del estado',
    'Gobernador': 'Gobierno del estado', 'Gobierno': 'Gobierno del estado',
    'Hacienda': 'Gobierno del estado', 'Oficialía Mayor': 'Gobierno del estado',
    'Oficialía_Mayor': 'Gobierno del estado', 'PGJEH': 'Gobierno del estado',
    'Planeación': 'Gobierno del estado', 'Saderh': 'Gobierno del estado',
    'Sedagroh': 'Gobierno del estado', 'Sedeco': 'Gobierno del estado',
    'Sedeso': 'Gobierno del estado', 'Segobh': 'Gobierno del estado',
    'Semarnath': 'Gobierno del estado', 'Semot': 'Gobierno del estado',
    'SEPH': 'Gobierno del estado', 'SIPDUS': 'Gobierno del estado',
    'Sopot': 'Gobierno del estado', 'SSH': 'Gobierno del estado',
    'SSPH': 'Gobierno del estado', 'Turismo': 'Gobierno del estado',
    'Upe': 'Gobierno del estado', 'STPSH': 'Gobierno del estado'
}

# --- CONFIGURACIÓN DE CAMPOS ---
COL_MEDIO = "Radiodifusora"
COL_AUTOR = "Autor"
COL_PROGRAMA = "Programa"
COL_FECHA = "Fecha"
COL_TITULO = "Titular"
COL_DEP = "Dependencia"
COL_ORG = "Organismo"
COL_TEMA = "Tema"
COL_MUN = "Municipio"
COL_EST = "Estatus"

root = Tk()
root.withdraw()
root.attributes('-topmost', True)

# 1. Seleccionar archivo de noticias
ruta_archivo = filedialog.askopenfilename(
    title="Selecciona el archivo de noticias (Nidia...)",
    filetypes=[("Archivos CSV", "*.csv")]
)

if not ruta_archivo:
    print("⚠️ No se seleccionó archivo.")
else:
    # 2. Seleccionar carpeta de destino
    ruta_salida = filedialog.askdirectory(title="¿Dónde quieres guardar los documentos?")

    if not ruta_salida:
        print("⚠️ No se seleccionó carpeta.")
    else:
        try:
            # Lectura con codificación para acentos (utf-8-sig es la más compatible)
            df = pd.read_csv(ruta_archivo, sep=None, engine='python', encoding='utf-8-sig', on_bad_lines='skip')
            df = df.dropna(how='all')
            
            # Normalizar nombres de columnas
            df.columns = [str(c).strip().capitalize() for c in df.columns]
            
            # Convertir fecha
            df[COL_FECHA] = pd.to_datetime(df[COL_FECHA], dayfirst=True, errors="coerce")

            archivos_creados = 0
            
            for medio, grupo in df.groupby(COL_MEDIO):
                if pd.isna(medio) or str(medio).strip() == "": continue
                    
                doc = Document()
                section = doc.sections[0]
                section.orientation = WD_ORIENT.LANDSCAPE
                section.page_width, section.page_height = section.page_height, section.page_width
                section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(1.27)

                table = doc.add_table(rows=2, cols=1)
                table.style = 'Table Grid'

                # Encabezado: Programa
                val_prog = str(grupo.iloc[0].get(COL_PROGRAMA, "S/P")).upper()
                p_header = table.cell(0, 0).paragraphs[0]
                run_h = p_header.add_run(val_prog)
                run_h.bold = True
                run_h.font.name = 'Arial'; run_h.font.size = Pt(12)
                p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER

                # Fila 2: Datos
                val_autor = str(grupo.iloc[0].get(COL_AUTOR, "S/A"))
                f_dt = grupo.iloc[0][COL_FECHA]
                f_str = f_dt.strftime("%d/%m/%Y") if pd.notnull(f_dt) else ""
                
                cell_body = table.cell(1, 0)
                p_intro = cell_body.paragraphs[0]
                run_intro = p_intro.add_run(f"{val_prog} - {val_autor} {f_str}: Impactos:")
                run_intro.font.name = 'Arial'; run_intro.font.size = Pt(10)

                # Clasificación de noticias
                impactos_gob = []
                impactos_dest = []
                negativas = []

                contador = 1
                for _, row in grupo.iterrows():
                    def clean(val): return str(val).strip() if pd.notnull(val) and str(val).lower() != 'nan' else ""
                    
                    dep = clean(row.get(COL_DEP))
                    org = clean(row.get(COL_ORG))
                    tema = clean(row.get(COL_TEMA))
                    mun = clean(row.get(COL_MUN))
                    tit = clean(row.get(COL_TITULO))
                    est = clean(row.get(COL_EST)).lower()

                    partes = [p for p in [dep, org, tema, mun] if p and p.lower() not in ['nan', 'n/a']]
                    encabezado = f"{' - '.join(partes)} - "
                    
                    # Búsqueda automática en el diccionario interno
                    tipo = CLASIFICACION_INTERNA.get(dep, "Gobierno del estado")

                    if "negativa" in est:
                        negativas.append((encabezado, tit))
                    elif tipo == "Destacadas":
                        impactos_dest.append((encabezado, tit))
                    else:
                        impactos_gob.append((encabezado, tit))
                    
                    contador += 1

                # Función para escribir bloques en el Word
                def escribir(lista, titulo=None):
                    if not lista: return
                    if titulo:
                        p = cell_body.add_paragraph()
                        r = p.add_run(titulo)
                        r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(10)
                    
                    for enc, t in lista:
                        p = cell_body.add_paragraph()
                        p.paragraph_format.space_after = Pt(0)
                        r_e = p.add_run(enc); r_e.bold = True
                        r_e.font.name = 'Arial'; r_e.font.size = Pt(10)
                        r_t = p.add_run(t); r_t.font.name = 'Arial'; r_t.font.size = Pt(10)

                # Orden de escritura
                escribir(impactos_gob)
                if impactos_dest: cell_body.add_paragraph()
                escribir(impactos_dest, "Destacadas:")
                if negativas: cell_body.add_paragraph()
                escribir(negativas, "Negativas:")

                # Guardar archivo
                f_sufijo = f_dt.strftime("%d%m%y") if pd.notnull(f_dt) else "000000"
                medio_nombre = "".join([c for c in str(medio) if c.isalnum() or c==' ']).strip()
                nombre_archivo = f"Resumen_{medio_nombre}_{f_sufijo}.docx"
                
                doc.save(os.path.join(ruta_salida, nombre_archivo))
                archivos_creados += 1

            messagebox.showinfo("Proceso automático listo", f"Se generaron {archivos_creados} archivos clasificados.")

        except Exception as e:
            messagebox.showerror("Error", f"Ocurrió un error: {e}")